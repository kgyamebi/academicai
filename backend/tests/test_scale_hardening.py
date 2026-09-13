from pathlib import Path

from app.core.pagination import decode_cursor, encode_cursor
from app.workers.queue import queue_depth, registered_workers, reset_redis_client


def test_ready_uses_cached_worker_count_not_direct_worker_all():
    source = Path(__file__).resolve().parents[1] / "app" / "api" / "v1" / "health.py"
    text = source.read_text(encoding="utf-8")
    assert "registered_workers()" in text
    assert "from rq" not in text
    assert "queue_meta=True" in text

def test_report_get_does_not_eager_load_all_findings():
    source = Path(__file__).resolve().parents[1] / "app" / "api" / "v1" / "reports.py"
    text = source.read_text(encoding="utf-8")
    assert "selectinload(AnalysisReport.findings)" not in text
    assert "selectinload(AnalysisReport.scores)" in text


def test_assignment_list_uses_direct_count():
    source = Path(__file__).resolve().parents[1] / "app" / "api" / "v1" / "assignments.py"
    text = source.read_text(encoding="utf-8")
    assert "func.count(Assignment.id)" in text
    assert "select_from(q.subquery())" not in text


def test_gzip_and_statement_cache_are_configured():
    main = (Path(__file__).resolve().parents[1] / "app" / "main.py").read_text(encoding="utf-8")
    session = (Path(__file__).resolve().parents[1] / "app" / "db" / "session.py").read_text(encoding="utf-8")
    assert "GZipMiddleware" in main
    assert "query_cache_size" in session
    assert "pool_use_lifo" in session


def test_queue_helpers_return_ints():
    reset_redis_client()
    workers = registered_workers()
    depth = queue_depth()
    assert isinstance(workers, int)
    assert isinstance(depth, int)
    assert workers >= 0
    assert depth >= -1


def test_assignment_list_supports_keyset_cursor():
    source = Path(__file__).resolve().parents[1] / "app" / "api" / "v1" / "assignments.py"
    text = source.read_text(encoding="utf-8")
    assert "keyset_before" in text
    assert "next_cursor" in text
    assert "Assignment.id.desc()" in text


def test_report_findings_support_keyset_cursor():
    source = Path(__file__).resolve().parents[1] / "app" / "api" / "v1" / "reports.py"
    text = source.read_text(encoding="utf-8")
    assert "keyset_after" in text
    assert "next_cursor" in text


def test_cursor_roundtrip_and_reject_garbage():
    from datetime import UTC, datetime
    from uuid import uuid4

    row_id = uuid4()
    ts = datetime.now(UTC)
    token = encode_cursor(ts, row_id)
    decoded_ts, decoded_id = decode_cursor(token)
    assert decoded_id == row_id
    assert decoded_ts.isoformat() == ts.isoformat()
    from fastapi import HTTPException

    try:
        decode_cursor("not-a-cursor")
        raise AssertionError("expected HTTPException")
    except HTTPException as exc:
        assert exc.status_code == 400


def test_assignment_cursor_pages_without_overlap(client):
    from uuid import uuid4

    email = f"scale-{uuid4().hex[:10]}@example.com"
    signup = client.post(
        "/api/auth/register",
        json={"email": email, "password": "password12", "full_name": "Scale"},
    )
    assert signup.status_code == 200, signup.text
    token = signup.json()["access_token"]
    headers = {"Authorization": f"Bearer {token}"}
    created = []
    for i in range(3):
        resp = client.post(
            "/api/assignments",
            json={"title": f"Scale {i}", "question": "Evaluate globalization."},
            headers=headers,
        )
        assert resp.status_code in {200, 201}, resp.text
        created.append(resp.json()["id"])
    first = client.get("/api/assignments", params={"page_size": 2}, headers=headers)
    assert first.status_code == 200
    body = first.json()
    assert "next_cursor" in body
    ids_first = [item["id"] for item in body["items"]]
    assert len(ids_first) == 2
    cursor = body["next_cursor"]
    assert cursor
    second = client.get("/api/assignments", params={"page_size": 2, "cursor": cursor}, headers=headers)
    assert second.status_code == 200
    ids_second = [item["id"] for item in second.json()["items"]]
    assert set(ids_first).isdisjoint(set(ids_second))
    bad = client.get("/api/assignments", params={"cursor": "%%%"}, headers=headers)
    assert bad.status_code == 400


def test_citations_list_batches_verifications():
    source = Path(__file__).resolve().parents[1] / "app" / "api" / "v1" / "citations.py"
    text = source.read_text(encoding="utf-8")
    assert "_latest_verifications" in text
    assert "group_by(SourceVerification.reference_id)" in text
    assert "_reference_payload(db, r)" not in text


def test_dashboard_loads_reports_once():
    source = Path(__file__).resolve().parents[1] / "app" / "api" / "v1" / "dashboard.py"
    text = source.read_text(encoding="utf-8")
    assert text.count("select(AnalysisReport)") == 1
    assert "limit(12)" in text
    assert "reports[:8]" in text
    assert "select(AnalysisReport.overall_score)" not in text


def test_public_faq_cache_hits_on_second_read(client):
    from app.core.app_cache import cache_clear
    from app.core.metrics import reset_for_tests, snapshot

    cache_clear()
    reset_for_tests()
    first = client.get("/api/public/faqs")
    second = client.get("/api/public/faqs")
    assert first.status_code == 200
    assert second.status_code == 200
    assert first.json() == second.json()
    counts = snapshot()
    assert counts.get("cache.meta.hit", 0) >= 1
    assert counts.get("cache.meta.miss", 0) >= 1


def test_verification_index_migration_exists():
    alembic = (
        Path(__file__).resolve().parents[1] / "alembic" / "versions" / "011_verification_index.py"
    )
    text = alembic.read_text(encoding="utf-8")
    assert "ix_source_verifications_reference_created" in text
    assert '["reference_id", "created_at"]' in text


def test_document_extract_size_tiers_write_evidence():
    import json
    import time
    import tracemalloc

    from app.services.documents.extractor import extract_document

    extract_document(b"Warm up the extractor so import cost is not in the samples.", ".txt")
    paragraph = (
        "Globalization has mixed effects on developing economies because trade rules "
        "differ across regions and institutions remain uneven.\n\n"
    )
    txt_tiers = {
        "small": paragraph * 8,
        "medium": paragraph * 80,
        "large": paragraph * 400,
        "very_large": paragraph * 1200,
    }
    txt_results = {}
    for name, text in txt_tiers.items():
        tracemalloc.start()
        start = time.perf_counter()
        extracted = extract_document(text.encode(), ".txt")
        elapsed_ms = (time.perf_counter() - start) * 1000
        _current, peak = tracemalloc.get_traced_memory()
        tracemalloc.stop()
        txt_results[name] = {
            "chars": len(text),
            "words": extracted.word_count,
            "paragraphs": extracted.paragraph_count,
            "latency_ms": round(elapsed_ms, 3),
            "peak_traced_bytes": peak,
        }
        assert extracted.word_count > 0
    assert txt_results["small"]["latency_ms"] < 2000
    assert txt_results["very_large"]["latency_ms"] < 15000

    pdf_results = _time_pdf_tiers()
    docx_results = _time_docx_tiers()
    out = Path(__file__).resolve().parents[2] / "ops" / "cert_document_extract.json"
    out.write_text(
        json.dumps(
            {
                "measured_at": time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime()),
                "allocator": "tracemalloc",
                "request_thread": "pdf_docx_extraction_moved_to_worker_queue",
                "txt": txt_results,
                "pdf": pdf_results,
                "docx": docx_results,
                "not_run": ["parallel_parse"],
            },
            indent=2,
        ),
        encoding="utf-8",
    )
    assert out.is_file()


def _time_pdf_tiers() -> dict:
    import time
    import tracemalloc

    import fitz

    from app.services.documents.extractor import extract_document

    page_counts = {"small": 1, "medium": 8, "large": 24, "very_large": 60}
    body = "Globalization has mixed effects on developing economies. " * 40
    results = {}
    warmup = fitz.open()
    page = warmup.new_page()
    page.insert_text((72, 72), body[:500])
    extract_document(warmup.tobytes(), ".pdf")
    warmup.close()
    for name, pages in page_counts.items():
        doc = fitz.open()
        for _ in range(pages):
            p = doc.new_page()
            p.insert_text((72, 72), body[:1500])
        blob = doc.tobytes()
        doc.close()
        tracemalloc.start()
        start = time.perf_counter()
        extracted = extract_document(blob, ".pdf")
        elapsed_ms = (time.perf_counter() - start) * 1000
        _current, peak = tracemalloc.get_traced_memory()
        tracemalloc.stop()
        results[name] = {
            "pages": pages,
            "bytes": len(blob),
            "words": extracted.word_count,
            "latency_ms": round(elapsed_ms, 3),
            "peak_traced_bytes": peak,
        }
        assert extracted.word_count > 0
    return results


def _time_docx_tiers() -> dict:
    import io
    import time
    import tracemalloc

    from docx import Document as DocxDocument

    from app.services.documents.extractor import extract_document

    paragraph_counts = {"small": 8, "medium": 80, "large": 400, "very_large": 1200}
    sentence = "Globalization has mixed effects on developing economies because institutions differ."
    results = {}
    warm = DocxDocument()
    warm.add_paragraph(sentence)
    buf = io.BytesIO()
    warm.save(buf)
    extract_document(buf.getvalue(), ".docx")
    for name, count in paragraph_counts.items():
        doc = DocxDocument()
        for _ in range(count):
            doc.add_paragraph(sentence)
        payload = io.BytesIO()
        doc.save(payload)
        blob = payload.getvalue()
        tracemalloc.start()
        start = time.perf_counter()
        extracted = extract_document(blob, ".docx")
        elapsed_ms = (time.perf_counter() - start) * 1000
        _current, peak = tracemalloc.get_traced_memory()
        tracemalloc.stop()
        results[name] = {
            "paragraphs_in": count,
            "bytes": len(blob),
            "words": extracted.word_count,
            "paragraphs": extracted.paragraph_count,
            "latency_ms": round(elapsed_ms, 3),
            "peak_traced_bytes": peak,
        }
        assert extracted.word_count > 0
    return results
