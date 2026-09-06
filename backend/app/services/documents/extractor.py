from __future__ import annotations

import io
import re
from dataclasses import dataclass, field

from app.config import get_settings
from app.services.documents.validation import DocumentSecurityError


@dataclass
class ExtractedParagraph:
    index: int
    text: str
    is_heading: bool = False
    heading_level: int = 0
    char_start: int = 0
    char_end: int = 0
    word_count: int = 0


@dataclass
class ExtractedSection:
    heading: str
    section_type: str
    start_paragraph: int
    end_paragraph: int
    sort_order: int


@dataclass
class ExtractedDocument:
    text: str
    normalized_text: str
    paragraphs: list[ExtractedParagraph] = field(default_factory=list)
    sections: list[ExtractedSection] = field(default_factory=list)
    word_count: int = 0
    word_count_excl_references: int = 0
    word_count_excl_headings: int = 0
    paragraph_count: int = 0
    sentence_count: int = 0
    page_count: int = 0
    language: str | None = None


HEADING_HINTS = re.compile(
    r"^(abstract|introduction|background|literature review|methodology|methods|"
    r"results|discussion|analysis|conclusion|conclusions|references|bibliography|"
    r"works cited|appendix|acknowledgements|counterargument|findings)\b",
    re.I,
)


def extract_document(content: bytes, extension: str) -> ExtractedDocument:
    settings = get_settings()
    if extension == ".pdf":
        raw, pages = _extract_pdf(content)
    elif extension == ".docx":
        raw, pages = _extract_docx(content)
    elif extension in {".txt", ".md"}:
        raw, pages = content.decode("utf-8", errors="replace"), 1
    else:
        raise DocumentSecurityError("Unsupported file type.")

    if len(raw) > settings.max_extracted_chars:
        raise DocumentSecurityError("The extracted text is too long for this plan and safety limit.")

    normalized = _normalize(raw)
    paragraphs = _split_paragraphs(normalized)
    sections = _detect_sections(paragraphs)
    words = _words(normalized)
    ref_start = _reference_start(paragraphs)
    body_text = " ".join(p.text for p in paragraphs if ref_start is None or p.index < ref_start)
    heading_excl = " ".join(p.text for p in paragraphs if not p.is_heading)
    sentences = _sentences(normalized)
    language = _detect_language(normalized)

    return ExtractedDocument(
        text=raw,
        normalized_text=normalized,
        paragraphs=paragraphs,
        sections=sections,
        word_count=len(words),
        word_count_excl_references=len(_words(body_text)),
        word_count_excl_headings=len(_words(heading_excl)),
        paragraph_count=len([p for p in paragraphs if not p.is_heading]),
        sentence_count=len(sentences),
        page_count=pages,
        language=language,
    )


def _extract_pdf(content: bytes) -> tuple[str, int]:
    import fitz

    settings = get_settings()
    try:
        doc = fitz.open(stream=content, filetype="pdf")
    except Exception as exc:
        raise DocumentSecurityError("We could not read this PDF. It may be damaged or password-protected.") from exc
    if doc.page_count > settings.max_pdf_pages:
        doc.close()
        raise DocumentSecurityError(f"This PDF has too many pages. Maximum is {settings.max_pdf_pages}.")
    if doc.is_encrypted:
        doc.close()
        raise DocumentSecurityError("Password-protected PDFs are not supported.")
    parts: list[str] = []
    for page in doc:
        parts.append(page.get_text("text") or "")
    pages = doc.page_count
    doc.close()
    return "\n\n".join(parts), pages


def _extract_docx(content: bytes) -> tuple[str, int]:
    from docx import Document as DocxDocument

    try:
        doc = DocxDocument(io.BytesIO(content))
    except Exception as exc:
        raise DocumentSecurityError("We could not read this Word document.") from exc
    blocks: list[str] = []
    for para in doc.paragraphs:
        style = (para.style.name if para.style is not None else "") or ""
        text = para.text.strip()
        if not text:
            continue
        if style.lower().startswith("heading"):
            blocks.append(text)
        else:
            blocks.append(text)
    return "\n\n".join(blocks), max(1, len(doc.paragraphs) // 12)


def _normalize(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _is_heading(text: str) -> bool:
    stripped = text.strip()
    if len(stripped) > 90:
        return False
    if HEADING_HINTS.match(stripped):
        return True
    if stripped.isupper() and 3 <= len(stripped.split()) <= 10:
        return True
    if re.match(r"^\d+(\.\d+)*\s+\S+", stripped) and len(stripped.split()) <= 12:
        return True
    return False


def _heading_level(text: str) -> int:
    m = re.match(r"^(\d+)(\.\d+)*\s+", text.strip())
    if m:
        return text.strip().count(".") + 1
    if HEADING_HINTS.match(text.strip()):
        return 1
    return 1


def _split_paragraphs(text: str) -> list[ExtractedParagraph]:
    chunks = [c.strip() for c in re.split(r"\n\s*\n", text) if c.strip()]
    paragraphs: list[ExtractedParagraph] = []
    cursor = 0
    for i, chunk in enumerate(chunks):
        start = text.find(chunk, cursor)
        end = start + len(chunk) if start >= 0 else cursor + len(chunk)
        cursor = end
        heading = _is_heading(chunk)
        paragraphs.append(
            ExtractedParagraph(
                index=i,
                text=chunk,
                is_heading=heading,
                heading_level=_heading_level(chunk) if heading else 0,
                char_start=max(start, 0),
                char_end=end,
                word_count=len(_words(chunk)),
            )
        )
    return paragraphs


def _section_type(heading: str) -> str:
    h = heading.lower()
    mapping = [
        ("abstract", "abstract"),
        ("introduction", "introduction"),
        ("background", "background"),
        ("literature", "literature"),
        ("method", "methods"),
        ("result", "results"),
        ("discussion", "discussion"),
        ("analysis", "analysis"),
        ("counter", "counterargument"),
        ("conclusion", "conclusion"),
        ("reference", "references"),
        ("bibliograph", "references"),
        ("works cited", "references"),
    ]
    for needle, value in mapping:
        if needle in h:
            return value
    return "body"


def _detect_sections(paragraphs: list[ExtractedParagraph]) -> list[ExtractedSection]:
    headings = [(i, p) for i, p in enumerate(paragraphs) if p.is_heading]
    if not headings:
        n = len(paragraphs)
        if n == 0:
            return []
        intro_end = max(1, n // 5)
        concl_start = max(intro_end, n - max(1, n // 6))
        return [
            ExtractedSection("Introduction", "introduction", 0, intro_end - 1, 0),
            ExtractedSection("Body", "body", intro_end, concl_start - 1, 1),
            ExtractedSection("Conclusion", "conclusion", concl_start, n - 1, 2),
        ]
    sections: list[ExtractedSection] = []
    for sort, (idx, para) in enumerate(headings):
        end = headings[sort + 1][0] - 1 if sort + 1 < len(headings) else len(paragraphs) - 1
        sections.append(
            ExtractedSection(
                heading=para.text[:120],
                section_type=_section_type(para.text),
                start_paragraph=idx,
                end_paragraph=end,
                sort_order=sort,
            )
        )
    return sections


def _reference_start(paragraphs: list[ExtractedParagraph]) -> int | None:
    for p in paragraphs:
        if p.is_heading and _section_type(p.text) == "references":
            return p.index
    return None


def _words(text: str) -> list[str]:
    return re.findall(r"[A-Za-z0-9']+", text)


def _sentences(text: str) -> list[str]:
    parts = re.split(r"(?<=[.!?])\s+", text)
    return [p.strip() for p in parts if p.strip()]


def _detect_language(text: str) -> str | None:
    try:
        from langdetect import detect

        sample = text[:4000]
        if len(sample) < 40:
            return None
        return detect(sample)
    except Exception:
        return None
